from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent
INPUT_XLSX = PROJECT_ROOT / 'public' / 'data' / 'NBArrwstoi History.xlsx'
OUTPUT_JSON = PROJECT_ROOT / 'public' / 'data' / 'history-data.json'

TEAM_ALIASES = {
    'Samarina Dudenbros2': 'Samarina Dudenbros',
    'Karafli SaiTa': 'Karafli Saita',
}

HEADER_MAP = {
    'Year': 'year',
    'Regular/\nPlayoffs': 'phase',
    'Period': 'period',
    'Match No': 'matchNo',
    'Team': 'team',
    'Manager': 'manager',
    'Games Won': 'gamesWon',
    'W': 'wins',
    'L': 'losses',
    'D': 'ties',
    'vs Opponent': 'opponent',
    'vs Manager': 'opponentManager',
    'FGM': 'fgm',
    'FG%': 'fgPct',
    '3PM': 'threePm',
    '3P%': 'threePct',
    'FT%': 'ftPct',
    'PTS': 'pts',
    'OREB': 'oreb',
    'DREB': 'dreb',
    'AST': 'ast',
    'STL': 'stl',
    'BLK': 'blk',
    'TO': 'to',
    'A/TO': 'ato',
}

PERCENT_KEYS = {'fgPct', 'threePct', 'ftPct'}
NUMERIC_KEYS = {
    'year', 'period', 'matchNo', 'gamesWon', 'wins', 'losses', 'ties', 'fgm', 'threePm',
    'pts', 'oreb', 'dreb', 'ast', 'stl', 'blk', 'to', 'ato'
}
FLOAT_KEYS = PERCENT_KEYS | {'gamesWon', 'wins', 'losses', 'ties', 'fgm', 'threePm', 'pts', 'oreb', 'dreb', 'ast', 'stl', 'blk', 'to', 'ato'}


def s(value: Any) -> str:
    return str(value or '').replace('\r', '').strip()


def canonical_team_name(value: Any) -> str:
    raw = s(value)
    return TEAM_ALIASES.get(raw, raw)


def normalize_phase(value: Any) -> str:
    text = s(value).lower()
    if 'playoff' in text:
        return 'Playoffs'
    if 'regular' in text:
        return 'Regular'
    return s(value)


def coerce_number(value: Any) -> Any:
    if value is None or value == '':
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            return int(value)
        return value

    text = s(value).replace(',', '')
    if not text:
        return None
    try:
        number = float(text)
        return int(number) if number.is_integer() else number
    except ValueError:
        return value


def load_rows() -> list[dict[str, Any]]:
    wb = load_workbook(INPUT_XLSX, data_only=True)
    ws = wb['Data']
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    rows: list[dict[str, Any]] = []

    for row_idx in range(2, ws.max_row + 1):
        row: dict[str, Any] = {}
        has_content = False
        for col_idx, header in enumerate(headers, start=1):
            key = HEADER_MAP.get(header)
            if not key:
                continue
            value = ws.cell(row_idx, col_idx).value
            if value not in (None, ''):
                has_content = True
            row[key] = value

        if not has_content:
            continue

        row['team'] = canonical_team_name(row.get('team'))
        row['opponent'] = canonical_team_name(row.get('opponent'))
        row['manager'] = s(row.get('manager'))
        row['opponentManager'] = s(row.get('opponentManager'))
        row['phase'] = normalize_phase(row.get('phase'))

        for key, value in list(row.items()):
            if key in NUMERIC_KEYS or key in PERCENT_KEYS:
                row[key] = coerce_number(value)
            elif key not in {'team', 'manager', 'opponent', 'opponentManager', 'phase'}:
                row[key] = coerce_number(value)

        if not row.get('year') or not row.get('team') or not row.get('opponent'):
            continue

        row['rowId'] = f"{row['year']}-{row['phase']}-{row['period']}-{row['matchNo']}-{slug(row['team'])}-{slug(row['opponent'])}"
        rows.append(row)

    return rows


def slug(value: str) -> str:
    value = s(value).lower()
    value = re.sub(r'[^a-z0-9]+', '-', value)
    return value.strip('-')


def main() -> None:
    rows = load_rows()
    payload = {
        'teamAliases': TEAM_ALIASES,
        'rowCount': len(rows),
        'years': sorted({row['year'] for row in rows}),
        'phases': sorted({row['phase'] for row in rows}),
        'teams': sorted({row['team'] for row in rows}),
        'rows': rows,
    }
    OUTPUT_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'Wrote {len(rows)} rows to {OUTPUT_JSON}')


if __name__ == '__main__':
    main()
